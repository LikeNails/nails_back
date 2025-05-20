import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Пути

project_root = r'E:\Projects\nails_back'
print(project_root)
output_file = 'Приложение_Код.docx'

# Расширения файлов, которые собираем
allowed_extensions = ['.ts', '.json']

# Создаем новый документ
doc = Document()

# Настройки стиля заголовков
style_heading = doc.styles['Heading 1']
font = style_heading.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)

# Настройки обычного текста (для кода)
style_normal = doc.styles['Normal']
font = style_normal.font
font.name = 'Courier New'
font.size = Pt(7)

# Счётчик для нумерации приложений
file_counter = 1

# Рекурсивный обход директории
for root, dirs, files in os.walk(project_root):
    for file in sorted(files):  # Чтобы файлы шли по алфавиту
        if any(file.endswith(ext) for ext in allowed_extensions):
            filepath = os.path.join(root, file)

            # Относительный путь или просто имя файла
            relative_path = os.path.relpath(filepath, project_root)

            # Добавляем заголовок
            heading = doc.add_heading(f'ПРИЛОЖЕНИЕ {file_counter}: Файл {relative_path}', level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

            # Принудительно задаём шрифт для заголовка
            run_heading = heading.runs[0]
            run_heading.font.name = 'Times New Roman'
            run_heading._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

            # Добавляем содержимое файла
            try:
                with open(filepath, 'r', encoding='utf-8') as f:
                    code = f.read()
            except Exception as e:
                code = f'[Ошибка чтения файла: {e}]'

            # Создаем абзац и сразу добавляем в него текст
            paragraph = doc.add_paragraph()
            run_paragraph = paragraph.add_run(code)
            run_paragraph.font.name = 'Courier New'
            run_paragraph._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
            run_paragraph.font.size = Pt(7)

            # Новый абзац для отступа между файлами
            doc.add_paragraph()

            # Увеличиваем счетчик
            file_counter += 1

# Сохраняем документ
doc.save(output_file)

print(f"\nДокумент сохранён как '{output_file}'")